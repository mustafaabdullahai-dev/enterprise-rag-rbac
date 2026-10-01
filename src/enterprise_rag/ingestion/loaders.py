"""Multi-format document loading: PDF, CSV/TSV, XLSX, DOCX, Markdown, text, JSON.

Each loader normalises its input into :class:`LoadedUnit` items that keep enough
structure for precise citations - a PDF page, a spreadsheet row, a markdown
section - instead of flattening everything into one blob of text. Citations
then point at ``[doc_id#chunk]`` and the UI can show *where* the answer came
from ("row 14 of the salary band spreadsheet").
"""

from __future__ import annotations

import csv
import hashlib
import io
import json
import re
from collections.abc import Iterable
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

MARKDOWN_SUFFIXES = {".md", ".markdown", ".txt", ".rst"}
RECORD_SUFFIXES = {".csv", ".tsv"}
PDF_SUFFIXES = {".pdf"}
DOCX_SUFFIXES = {".docx"}
EXCEL_SUFFIXES = {".xlsx", ".xlsm"}
JSON_SUFFIXES = {".json"}

SUPPORTED_SUFFIXES = (
    MARKDOWN_SUFFIXES
    | RECORD_SUFFIXES
    | PDF_SUFFIXES
    | DOCX_SUFFIXES
    | EXCEL_SUFFIXES
    | JSON_SUFFIXES
)


class LoaderError(RuntimeError):
    """Raised when a document cannot be read."""


@dataclass(frozen=True, slots=True)
class LoadedUnit:
    """One citable unit of a document (page, row, section, paragraph...)."""

    ref: str
    text: str
    metadata: dict[str, Any] = field(default_factory=dict)


@dataclass(frozen=True, slots=True)
class LoadedDocument:
    """A loaded document with its ordered units."""

    doc_id: str
    title: str
    source_path: str
    doc_format: str
    units: tuple[LoadedUnit, ...]

    @property
    def text(self) -> str:
        return "\n\n".join(unit.text for unit in self.units)


def slugify_doc_id(relative_path: str) -> str:
    """Deterministic, citation-safe document id derived from the file path."""
    stem = Path(relative_path).with_suffix("").as_posix()
    slug = re.sub(r"[^A-Za-z0-9]+", "-", stem).strip("-").lower()
    slug = re.sub(r"-{2,}", "-", slug)[:48].strip("-")
    if not slug:
        slug = "doc"
        slug += "-" + hashlib.sha1(relative_path.encode()).hexdigest()[:6]
    return slug


def load_document(path: Path, *, corpus_root: Path | None = None) -> LoadedDocument:
    """Load any supported file into structured units."""
    path = Path(path)
    if not path.exists():
        raise LoaderError(f"File not found: {path}")
    suffix = path.suffix.lower()

    try:
        relative = path.relative_to(corpus_root).as_posix() if corpus_root else path.name
    except ValueError:
        relative = path.name

    doc_id = slugify_doc_id(relative)
    title = _title_from(path)

    if suffix in PDF_SUFFIXES:
        units = _load_pdf(path)
    elif suffix in RECORD_SUFFIXES:
        units = _load_records(path, title)
    elif suffix in EXCEL_SUFFIXES:
        units = _load_excel(path, title)
    elif suffix in DOCX_SUFFIXES:
        units = _load_docx(path)
    elif suffix in MARKDOWN_SUFFIXES:
        units = _load_markdown(path)
    elif suffix in JSON_SUFFIXES:
        units = _load_json(path)
    else:
        raise LoaderError(f"Unsupported file type: {suffix}")

    units = tuple(unit for unit in units if unit.text.strip())
    if not units:
        raise LoaderError(f"No extractable text in {path.name}")
    return LoadedDocument(
        doc_id=doc_id,
        title=title,
        source_path=str(path),
        doc_format=suffix.lstrip("."),
        units=units,
    )


def iter_corpus_files(root: Path) -> Iterable[Path]:
    """Yield every supported file under ``root`` (sorted, skipping sidecars)."""
    root = Path(root)
    if root.is_file():
        yield root
        return
    for path in sorted(root.rglob("*")):
        if not path.is_file():
            continue
        if path.name.startswith("_") or path.name.startswith("."):
            continue
        if path.suffix.lower() in SUPPORTED_SUFFIXES:
            yield path


def _title_from(path: Path) -> str:
    return path.stem.replace("_", " ").replace("-", " ").strip().title()


# --------------------------------------------------------------------------- #
# Markdown / plain text - split on headings, keep a breadcrumb per section
# --------------------------------------------------------------------------- #
_HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")


def _load_markdown(path: Path) -> list[LoadedUnit]:
    lines = path.read_text(encoding="utf-8", errors="replace").splitlines()
    units: list[LoadedUnit] = []
    breadcrumb: list[str] = []
    buffer: list[str] = []
    counter = 0

    def flush() -> None:
        nonlocal buffer, counter
        body = "\n".join(buffer).strip()
        if body:
            counter += 1
            prefix = " > ".join(breadcrumb)
            text = f"{prefix}\n{body}" if prefix else body
            units.append(
                LoadedUnit(
                    ref=f"section:{counter}",
                    text=text,
                    metadata={"unit": "section", "section": prefix or "body"},
                )
            )
        buffer = []

    for line in lines:
        match = _HEADING_RE.match(line)
        if match:
            flush()
            level = len(match.group(1))
            heading = match.group(2).strip()
            breadcrumb = breadcrumb[: level - 1]
            breadcrumb.append(heading)
        else:
            buffer.append(line)
    flush()
    return units


# --------------------------------------------------------------------------- #
# CSV / TSV - one citable unit per row, plus a schema unit
# --------------------------------------------------------------------------- #
def _load_records(path: Path, title: str) -> list[LoadedUnit]:
    delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
    raw = path.read_text(encoding="utf-8-sig", errors="replace")
    reader = csv.DictReader(io.StringIO(raw), delimiter=delimiter)
    rows = list(reader)
    columns = [c for c in (reader.fieldnames or []) if c]
    units: list[LoadedUnit] = []

    if columns:
        units.append(
            LoadedUnit(
                ref="schema",
                text=(
                    f"{title} - tabular dataset with columns: "
                    + ", ".join(columns)
                    + f". Contains {len(rows)} rows."
                ),
                metadata={"unit": "schema", "columns": ", ".join(columns), "rows": len(rows)},
            )
        )

    for index, row in enumerate(rows, start=1):
        cells = [
            f"{key.strip()}: {str(value).strip()}" for key, value in row.items() if key and value
        ]
        units.append(
            LoadedUnit(
                ref=f"row:{index}",
                text=f"{title} - record {index} of {len(rows)}. " + " | ".join(cells),
                metadata={"unit": "row", "row_index": index, "record": " | ".join(cells)},
            )
        )
    return units


# --------------------------------------------------------------------------- #
# XLSX - same idea as CSV, per worksheet
# --------------------------------------------------------------------------- #
def _load_excel(path: Path, title: str) -> list[LoadedUnit]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - optional dependency
        raise LoaderError("openpyxl is required to read .xlsx files.") from exc

    workbook = load_workbook(filename=str(path), read_only=True, data_only=True)
    units: list[LoadedUnit] = []
    try:
        for sheet in workbook.worksheets:
            rows = [[cell for cell in row] for row in sheet.iter_rows(values_only=True)]
            rows = [
                row for row in rows if any(cell is not None and str(cell).strip() for cell in row)
            ]
            if not rows:
                continue
            header = [
                str(cell).strip() if cell is not None else f"col{i}"
                for i, cell in enumerate(rows[0])
            ]
            units.append(
                LoadedUnit(
                    ref=f"sheet:{sheet.title}:schema",
                    text=(
                        f"{title} - worksheet '{sheet.title}' with columns: "
                        + ", ".join(header)
                        + f". Contains {max(0, len(rows) - 1)} data rows."
                    ),
                    metadata={"unit": "schema", "sheet": sheet.title, "columns": ", ".join(header)},
                )
            )
            for index, row in enumerate(rows[1:], start=1):
                cells = [
                    f"{header[i] if i < len(header) else f'col{i}'}: {str(value).strip()}"
                    for i, value in enumerate(row)
                    if value is not None and str(value).strip()
                ]
                units.append(
                    LoadedUnit(
                        ref=f"sheet:{sheet.title}:row:{index}",
                        text=(
                            f"{title} - worksheet '{sheet.title}', row {index}. "
                            + " | ".join(cells)
                        ),
                        metadata={"unit": "row", "sheet": sheet.title, "row_index": index},
                    )
                )
    finally:
        workbook.close()
    return units


# --------------------------------------------------------------------------- #
# PDF - one unit per page, preserving page numbers for citations
# --------------------------------------------------------------------------- #
def _load_pdf(path: Path) -> list[LoadedUnit]:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - optional dependency
        raise LoaderError("pypdf is required to read PDF files.") from exc

    try:
        reader = PdfReader(str(path))
    except Exception as exc:  # pragma: no cover - malformed file
        raise LoaderError(f"Could not parse PDF {path.name}: {exc}") from exc

    units: list[LoadedUnit] = []
    for page_number, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        text = re.sub(r"[ \t]{2,}", " ", text)
        if not text:
            continue
        units.append(
            LoadedUnit(
                ref=f"page:{page_number}",
                text=text,
                metadata={"unit": "page", "page": page_number},
            )
        )
    return units


# --------------------------------------------------------------------------- #
# DOCX - heading-aware sections, tables rendered as pipe-delimited text
# --------------------------------------------------------------------------- #
def _load_docx(path: Path) -> list[LoadedUnit]:
    try:
        import docx  # type: ignore[import-untyped]
    except ImportError as exc:  # pragma: no cover - optional dependency
        raise LoaderError("python-docx is required to read .docx files.") from exc

    document = docx.Document(str(path))
    units: list[LoadedUnit] = []
    breadcrumb: list[str] = []
    buffer: list[str] = []
    counter = 0

    def flush() -> None:
        nonlocal buffer, counter
        body = "\n".join(buffer).strip()
        if body:
            counter += 1
            prefix = " > ".join(breadcrumb)
            text = f"{prefix}\n{body}" if prefix else body
            units.append(
                LoadedUnit(
                    ref=f"section:{counter}",
                    text=text,
                    metadata={"unit": "section", "section": prefix or "body"},
                )
            )
        buffer = []

    for paragraph in document.paragraphs:
        style = (paragraph.style.name or "").lower() if paragraph.style else ""
        text = paragraph.text.strip()
        if not text:
            continue
        if style.startswith("heading"):
            flush()
            level = 1
            match = re.search(r"(\d+)", style)
            if match:
                level = int(match.group(1))
            breadcrumb = breadcrumb[: level - 1]
            breadcrumb.append(text)
        else:
            buffer.append(text)
    flush()

    for table_index, table in enumerate(document.tables, start=1):
        rows = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                rows.append(" | ".join(cells))
        if rows:
            units.append(
                LoadedUnit(
                    ref=f"table:{table_index}",
                    text=f"Table {table_index}\n" + "\n".join(rows),
                    metadata={"unit": "table", "table_index": table_index},
                )
            )
    return units


# --------------------------------------------------------------------------- #
# JSON - either a list of records or a mapping of key -> value
# --------------------------------------------------------------------------- #
def _load_json(path: Path) -> list[LoadedUnit]:
    payload = json.loads(path.read_text(encoding="utf-8", errors="replace"))
    units: list[LoadedUnit] = []
    records = payload if isinstance(payload, list) else [payload]
    for index, record in enumerate(records, start=1):
        if isinstance(record, dict):
            cells = [f"{k}: {json.dumps(v, ensure_ascii=False)}" for k, v in record.items()]
        else:
            cells = [json.dumps(record, ensure_ascii=False)]
        units.append(
            LoadedUnit(
                ref=f"record:{index}",
                text=f"Record {index} of {len(records)}. " + " | ".join(cells),
                metadata={"unit": "record", "record_index": index},
            )
        )
    return units


__all__ = [
    "iter_corpus_files",
    "LoadedDocument",
    "LoadedUnit",
    "LoaderError",
    "load_document",
    "slugify_doc_id",
    "SUPPORTED_SUFFIXES",
]
